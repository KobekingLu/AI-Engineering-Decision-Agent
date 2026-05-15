"""Build a SQLite knowledge base from the historical bug review materials.

This script preserves provenance by storing:

- the source documents themselves
- every extracted raw bug row
- a canonical deduplicated bug table used by the decision agent

It is safe to rerun. The database tables are rebuilt from scratch each time.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sqlite3
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import fitz
import extract_msg
from openpyxl import load_workbook
from pptx import Presentation


SOURCE_KIND_PRIORITY = {
    "xlsx": 30,
    "csv": 25,
    "pdf": 18,
    "msg": 12,
    "pptx": 10,
    "ppt": 8,
}

HEADER_GROUPS = {
    "bug_id": ["bug id", "#", "issue id", "id"],
    "project": ["project", "案件"],
    "component": ["component"],
    "status": ["status"],
    "subject": ["subject", "title"],
    "description": ["說明", "description", "detail"],
    "rd_comment": ["rd comment", "rd command", "r&d solution for fix", "solution(root cause)", "solution (root cause)"],
    "author": ["author"],
    "assignee": ["assignee"],
    "feature_id": ["feature id"],
    "feature_id_detail": ["feature id detail"],
    "created_time": ["created time", "created"],
    "closed_time": ["closed"],
    "tracker": ["tracker"],
    "severity": ["severity"],
    "error_type": ["error type"],
    "issue_finder": ["issue finder"],
    "hw_version": ["hw version"],
    "fw_version": ["fw version"],
    "counts": ["counts"],
    "updated": ["updated"],
    "percent_done": ["% done", "percent done"],
}

DESCRIPTION_GROUPS = [
    ("說明", ["description", "說明", "detail"]),
    ("RD Comment", ["rd_comment", "rd comment", "rd command"]),
    ("R&D Solution for fix", ["r&d solution for fix"]),
    ("Solution(Root cause)", ["solution(root cause)", "solution (root cause)"]),
    ("Feature ID Detail", ["feature_id_detail"]),
]

BUG_ROW_REQUIRED_GROUPS = {"bug_id", "project", "status", "subject", "component"}


@dataclass(slots=True)
class SourceDocument:
    relative_path: str
    file_name: str
    file_extension: str
    source_kind: str
    quarter: str
    title: str
    extracted_text: str
    file_size: int
    sha256: str
    metadata: dict[str, Any]


@dataclass(slots=True)
class BugRecord:
    bug_id: str
    dedupe_key: str
    project: str
    component: str
    status: str
    subject: str
    description: str
    quarter: str
    source_sheet: str
    source_path: str
    source_kind: str
    source_row_number: int
    period_rank: int
    source_priority: int
    record_score: int
    raw_json: dict[str, Any]


def main(argv: list[str] | None = None) -> Path:
    parser = argparse.ArgumentParser(description="Build the local SQLite knowledge base.")
    parser.add_argument(
        "--repo-root",
        type=Path,
        default=Path(__file__).resolve().parents[1],
        help="Repository root. Defaults to the parent directory of this script.",
    )
    parser.add_argument(
        "--db-path",
        type=Path,
        default=None,
        help="Optional explicit SQLite output path.",
    )
    parser.add_argument(
        "--review-root",
        type=Path,
        default=None,
        help="Optional explicit BUG Review source root.",
    )
    parser.add_argument(
        "--csv-path",
        type=Path,
        default=None,
        help="Optional explicit merged CSV source path.",
    )
    args = parser.parse_args(argv)

    repo_root = args.repo_root.resolve()
    decision_root = repo_root / "decision_agent"
    knowledge_root = decision_root / "knowledge_base"
    review_root = (args.review_root or (knowledge_root / "BUG Review")).resolve()
    csv_path = (args.csv_path or (knowledge_root / "redmine_bug_summary_merged.csv")).resolve()
    db_path = (args.db_path or (repo_root / ".tmp" / "knowledge_base.generated.sqlite3")).resolve()

    documents, bug_records, stats = collect_knowledge_base_inputs(repo_root, review_root, csv_path)
    canonical_records = choose_canonical_bug_records(bug_records)
    build_database(db_path, documents, bug_records, canonical_records)

    summary = {
        "db_path": str(db_path),
        "source_document_count": len(documents),
        "raw_bug_record_count": len(bug_records),
        "canonical_bug_record_count": len(canonical_records),
        "stats": stats,
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return db_path


def collect_knowledge_base_inputs(
    repo_root: Path,
    review_root: Path,
    csv_path: Path,
) -> tuple[list[SourceDocument], list[BugRecord], dict[str, Any]]:
    documents: list[SourceDocument] = []
    bug_records: list[BugRecord] = []
    stats: dict[str, Any] = {
        "xlsx_files": 0,
        "msg_files": 0,
        "pptx_files": 0,
        "ppt_files": 0,
        "pdf_files": 0,
        "csv_rows": 0,
        "warnings": [],
    }

    csv_document, csv_records, csv_stats = _collect_csv_source(csv_path, repo_root)
    documents.append(csv_document)
    bug_records.extend(csv_records)
    stats["csv_rows"] = csv_stats["row_count"]

    for path in sorted(review_root.rglob("*")):
        if not path.is_file():
            continue
        suffix = path.suffix.lower()
        try:
            if suffix == ".xlsx":
                document, records, sheet_stats = _collect_xlsx_source(path, repo_root)
                stats["xlsx_files"] += 1
                stats.setdefault("xlsx_sheets", 0)
                stats["xlsx_sheets"] += sheet_stats["sheet_count"]
            elif suffix == ".msg":
                document = _collect_msg_source(path, repo_root)
                records = []
                stats["msg_files"] += 1
            elif suffix == ".pptx":
                document = _collect_pptx_source(path, repo_root)
                records = []
                stats["pptx_files"] += 1
            elif suffix == ".ppt":
                document = _collect_ppt_source(path, repo_root)
                records = []
                stats["ppt_files"] += 1
            elif suffix == ".pdf":
                document = _collect_pdf_source(path, repo_root)
                records = []
                stats["pdf_files"] += 1
            else:
                continue
        except Exception as exc:  # pragma: no cover - defensive build-time handling
            document = _build_error_source_document(path, repo_root, exc)
            records = []
            stats["warnings"].append(f"{path}: {exc}")
        documents.append(document)
        bug_records.extend(records)

    return documents, bug_records, stats


def build_database(
    db_path: Path,
    documents: list[SourceDocument],
    bug_records: list[BugRecord],
    canonical_records: list[BugRecord],
) -> None:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(db_path) as connection:
        connection.execute("PRAGMA foreign_keys = ON")
        connection.execute("PRAGMA journal_mode = OFF")
        connection.execute("PRAGMA synchronous = OFF")
        connection.execute("PRAGMA temp_store = MEMORY")
        _create_schema(connection)

        document_id_by_path = _insert_source_documents(connection, documents)
        _insert_bug_records(connection, bug_records, document_id_by_path)
        _insert_canonical_bug_records(connection, canonical_records, document_id_by_path)
        _create_indexes(connection)

        connection.commit()


def _create_schema(connection: sqlite3.Connection) -> None:
    connection.executescript(
        """
        DROP TABLE IF EXISTS canonical_bug_rows;
        DROP TABLE IF EXISTS source_bug_records;
        DROP TABLE IF EXISTS source_documents;

        CREATE TABLE source_documents (
            source_document_id INTEGER PRIMARY KEY AUTOINCREMENT,
            relative_path TEXT NOT NULL UNIQUE,
            file_name TEXT NOT NULL,
            file_extension TEXT NOT NULL,
            source_kind TEXT NOT NULL,
            quarter TEXT NOT NULL DEFAULT '',
            title TEXT NOT NULL DEFAULT '',
            extracted_text TEXT NOT NULL DEFAULT '',
            file_size INTEGER NOT NULL DEFAULT 0,
            sha256 TEXT NOT NULL DEFAULT '',
            metadata_json TEXT NOT NULL DEFAULT '{}',
            imported_at TEXT NOT NULL
        );

        CREATE TABLE source_bug_records (
            record_id INTEGER PRIMARY KEY AUTOINCREMENT,
            bug_id TEXT NOT NULL,
            dedupe_key TEXT NOT NULL,
            project TEXT NOT NULL DEFAULT '',
            component TEXT NOT NULL DEFAULT '',
            status TEXT NOT NULL DEFAULT '',
            subject TEXT NOT NULL DEFAULT '',
            description TEXT NOT NULL DEFAULT '',
            quarter TEXT NOT NULL DEFAULT '',
            source_sheet TEXT NOT NULL DEFAULT '',
            source_path TEXT NOT NULL DEFAULT '',
            source_kind TEXT NOT NULL DEFAULT '',
            source_document_id INTEGER,
            source_row_number INTEGER NOT NULL DEFAULT 0,
            period_rank INTEGER NOT NULL DEFAULT 0,
            source_priority INTEGER NOT NULL DEFAULT 0,
            record_score INTEGER NOT NULL DEFAULT 0,
            raw_json TEXT NOT NULL DEFAULT '{}',
            FOREIGN KEY(source_document_id) REFERENCES source_documents(source_document_id)
        );

        CREATE TABLE canonical_bug_rows (
            bug_id TEXT NOT NULL,
            dedupe_key TEXT NOT NULL,
            project TEXT NOT NULL DEFAULT '',
            component TEXT NOT NULL DEFAULT '',
            status TEXT NOT NULL DEFAULT '',
            subject TEXT NOT NULL DEFAULT '',
            description TEXT NOT NULL DEFAULT '',
            quarter TEXT NOT NULL DEFAULT '',
            source_sheet TEXT NOT NULL DEFAULT '',
            source_path TEXT NOT NULL DEFAULT '',
            source_kind TEXT NOT NULL DEFAULT '',
            source_document_id INTEGER,
            source_row_number INTEGER NOT NULL DEFAULT 0,
            period_rank INTEGER NOT NULL DEFAULT 0,
            source_priority INTEGER NOT NULL DEFAULT 0,
            record_score INTEGER NOT NULL DEFAULT 0,
            raw_json TEXT NOT NULL DEFAULT '{}',
            FOREIGN KEY(source_document_id) REFERENCES source_documents(source_document_id)
        );
        """
    )


def _insert_source_documents(
    connection: sqlite3.Connection,
    documents: list[SourceDocument],
) -> dict[str, int]:
    mapping: dict[str, int] = {}
    cursor = connection.cursor()
    for document in documents:
        cursor.execute(
            """
            INSERT INTO source_documents (
                relative_path,
                file_name,
                file_extension,
                source_kind,
                quarter,
                title,
                extracted_text,
                file_size,
                sha256,
                metadata_json,
                imported_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                document.relative_path,
                document.file_name,
                document.file_extension,
                document.source_kind,
                document.quarter,
                document.title,
                document.extracted_text,
                document.file_size,
                document.sha256,
                json.dumps(document.metadata, ensure_ascii=False, default=str),
                _utc_now(),
            ),
        )
        mapping[document.relative_path] = int(cursor.lastrowid)
    return mapping


def _insert_bug_records(
    connection: sqlite3.Connection,
    bug_records: list[BugRecord],
    document_id_by_path: dict[str, int],
) -> None:
    cursor = connection.cursor()
    for record in bug_records:
        cursor.execute(
            """
            INSERT INTO source_bug_records (
                bug_id,
                dedupe_key,
                project,
                component,
                status,
                subject,
                description,
                quarter,
                source_sheet,
                source_path,
                source_kind,
                source_document_id,
                source_row_number,
                period_rank,
                source_priority,
                record_score,
                raw_json
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                record.bug_id,
                record.dedupe_key,
                record.project,
                record.component,
                record.status,
                record.subject,
                record.description,
                record.quarter,
                record.source_sheet,
                record.source_path,
                record.source_kind,
                document_id_by_path.get(record.source_path),
                record.source_row_number,
                record.period_rank,
                record.source_priority,
                record.record_score,
                json.dumps(record.raw_json, ensure_ascii=False, default=str),
            ),
        )


def _insert_canonical_bug_records(
    connection: sqlite3.Connection,
    canonical_records: list[BugRecord],
    document_id_by_path: dict[str, int],
) -> None:
    cursor = connection.cursor()
    for record in canonical_records:
        cursor.execute(
            """
            INSERT INTO canonical_bug_rows (
                bug_id,
                dedupe_key,
                project,
                component,
                status,
                subject,
                description,
                quarter,
                source_sheet,
                source_path,
                source_kind,
                source_document_id,
                source_row_number,
                period_rank,
                source_priority,
                record_score,
                raw_json
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                record.bug_id,
                record.dedupe_key,
                record.project,
                record.component,
                record.status,
                record.subject,
                record.description,
                record.quarter,
                record.source_sheet,
                record.source_path,
                record.source_kind,
                document_id_by_path.get(record.source_path),
                record.source_row_number,
                record.period_rank,
                record.source_priority,
                record.record_score,
                json.dumps(record.raw_json, ensure_ascii=False, default=str),
            ),
        )


def _create_indexes(connection: sqlite3.Connection) -> None:
    connection.executescript(
        """
        CREATE INDEX idx_source_bug_records_bug_id ON source_bug_records(bug_id);
        CREATE INDEX idx_source_bug_records_dedupe_key ON source_bug_records(dedupe_key);
        CREATE INDEX idx_canonical_bug_rows_bug_id ON canonical_bug_rows(bug_id);
        CREATE INDEX idx_canonical_bug_rows_project ON canonical_bug_rows(project);
        CREATE INDEX idx_canonical_bug_rows_quarter ON canonical_bug_rows(quarter);
        CREATE INDEX idx_source_documents_relative_path ON source_documents(relative_path);
        """
    )


def choose_canonical_bug_records(bug_records: list[BugRecord]) -> list[BugRecord]:
    grouped: dict[str, list[BugRecord]] = defaultdict(list)
    for record in bug_records:
        grouped[record.dedupe_key].append(record)

    canonical_records: list[BugRecord] = []
    for records in grouped.values():
        canonical_records.append(max(records, key=_selection_key))

    canonical_records.sort(key=lambda record: (record.bug_id == "", record.bug_id.lower(), record.quarter, record.source_sheet))
    return canonical_records


def _selection_key(record: BugRecord) -> tuple[int, int, int, int, int, int]:
    completeness = sum(
        1
        for value in [record.bug_id, record.project, record.component, record.status, record.subject, record.description]
        if value
    )
    return (
        record.period_rank,
        record.source_priority,
        completeness,
        len(record.subject),
        len(record.description),
        -record.source_row_number,
        )


def _collect_csv_source(csv_path: Path, repo_root: Path) -> tuple[SourceDocument, list[BugRecord], dict[str, Any]]:
    rows: list[dict[str, str]] = []
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        rows = [row for row in reader]

    source_path = str(csv_path.resolve().relative_to(repo_root))
    document = SourceDocument(
        relative_path=source_path,
        file_name=csv_path.name,
        file_extension=csv_path.suffix.lower(),
        source_kind="csv",
        quarter=_infer_quarter_from_text(csv_path.stem),
        title="Merged Redmine bug summary",
        extracted_text=f"CSV source with {len(rows)} rows.",
        file_size=csv_path.stat().st_size,
        sha256=_sha256(csv_path),
        metadata={
            "row_count": len(rows),
            "columns": list(rows[0].keys()) if rows else [],
        },
    )

    bug_records = []
    for row_number, row in enumerate(rows, start=2):
        row_values = {
            _canonical_key_for_header(key): str(value or "").strip()
            for key, value in row.items()
            if key
        }
        bug_records.append(
            _build_bug_record_from_mapping(
                row_values=row_values,
                source_path=source_path,
                source_kind="csv",
                source_sheet=str(row.get("source_sheet", "") or ""),
                source_row_number=row_number,
                quarter=str(row.get("quarter", "") or ""),
                source_priority=SOURCE_KIND_PRIORITY["csv"],
                extra_description_parts=[],
                extra_raw_fields={},
            )
        )

    return document, bug_records, {"row_count": len(rows)}


def _collect_xlsx_source(xlsx_path: Path, repo_root: Path) -> tuple[SourceDocument, list[BugRecord], dict[str, Any]]:
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        issue_sheet_stats: list[dict[str, Any]] = []
        bug_records: list[BugRecord] = []
        workbook_quarter = _infer_quarter_from_path(xlsx_path)
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            header_row_number, header_map, canonical_header_map = _find_issue_sheet_header(sheet)
            if header_row_number is None:
                continue
            sheet_records, sheet_stats = _collect_xlsx_sheet_records(
                workbook_path=xlsx_path,
                repo_root=repo_root,
                sheet_name=sheet_name,
                sheet=sheet,
                header_row_number=header_row_number,
                header_map=header_map,
                canonical_header_map=canonical_header_map,
                workbook_quarter=workbook_quarter,
            )
            if not sheet_records:
                continue
            issue_sheet_stats.append(sheet_stats)
            bug_records.extend(sheet_records)

        document = SourceDocument(
            relative_path=str(xlsx_path.resolve().relative_to(repo_root)),
            file_name=xlsx_path.name,
            file_extension=xlsx_path.suffix.lower(),
            source_kind="xlsx",
            quarter=workbook_quarter,
            title=issue_sheet_stats[0]["sheet_name"] if issue_sheet_stats else xlsx_path.stem,
            extracted_text=_summarize_xlsx_document(issue_sheet_stats),
            file_size=xlsx_path.stat().st_size,
            sha256=_sha256(xlsx_path),
            metadata={
                "sheet_count": len(workbook.sheetnames),
                "issue_sheet_count": len(issue_sheet_stats),
                "issue_sheets": issue_sheet_stats,
            },
        )
        return document, bug_records, {"sheet_count": len(issue_sheet_stats)}
    finally:
        workbook.close()


def _collect_xlsx_sheet_records(
    *,
    workbook_path: Path,
    repo_root: Path,
    sheet_name: str,
    sheet: Any,
    header_row_number: int,
    header_map: dict[int, str],
    canonical_header_map: dict[int, str],
    workbook_quarter: str,
) -> tuple[list[BugRecord], dict[str, Any]]:
    records: list[BugRecord] = []
    row_count = 0
    sample_subject = ""
    sheet_quarter = _infer_quarter_from_text(sheet_name) or workbook_quarter

    for row_number, row in enumerate(sheet.iter_rows(min_row=header_row_number + 1, values_only=True), start=header_row_number + 1):
        if not any(cell not in (None, "") for cell in row):
            continue
        row_count += 1
        row_mapping = _row_to_mapping(row, header_map)
        row_values = _row_to_canonical_values(row_mapping, header_map)
        bug_record = _build_bug_record_from_mapping(
            row_values=row_values,
            source_path=str(workbook_path.resolve().relative_to(repo_root)),
            source_kind="xlsx",
            source_sheet=sheet_name,
            source_row_number=row_number,
            quarter=sheet_quarter,
            source_priority=SOURCE_KIND_PRIORITY["xlsx"],
            extra_description_parts=_collect_description_parts(row_values),
            extra_raw_fields={
                "sheet_name": sheet_name,
                "workbook_path": str(workbook_path),
                "source_row_number": row_number,
                "header_map": canonical_header_map,
            },
        )
        if _is_header_or_summary_row(bug_record):
            continue
        if not bug_record.bug_id and not bug_record.subject and not bug_record.description:
            continue
        if not sample_subject and bug_record.subject:
            sample_subject = bug_record.subject
        records.append(bug_record)

    stats = {
        "sheet_name": sheet_name,
        "row_count": row_count,
        "sample_subject": sample_subject,
        "quarter": sheet_quarter,
    }
    return records, stats


def _collect_msg_source(msg_path: Path, repo_root: Path) -> SourceDocument:
    message = extract_msg.Message(str(msg_path))
    try:
        subject = str(getattr(message, "subject", "") or "").strip()
        sender = str(getattr(message, "sender", "") or "").strip()
        to = str(getattr(message, "to", "") or "").strip()
        sent_at = str(getattr(message, "date", "") or "").strip()
        attachments = getattr(message, "attachments", []) or []
        body = str(getattr(message, "body", "") or "")
        extracted_text = _truncate_text(body, 20000)
        metadata = {
            "subject": subject,
            "sender": sender,
            "to": to,
            "date": sent_at,
            "attachment_count": len(attachments),
            "attachment_names": [str(getattr(item, "longFilename", "") or getattr(item, "shortFilename", "") or "") for item in attachments],
            "extraction_status": "ok",
        }
        title = subject or msg_path.stem
        return SourceDocument(
            relative_path=str(msg_path.resolve().relative_to(repo_root)),
            file_name=msg_path.name,
            file_extension=msg_path.suffix.lower(),
            source_kind="msg",
            quarter=_infer_quarter_from_path(msg_path),
            title=title,
            extracted_text=extracted_text,
            file_size=msg_path.stat().st_size,
            sha256=_sha256(msg_path),
            metadata=metadata,
        )
    finally:
        close = getattr(message, "close", None)
        if callable(close):
            close()


def _collect_pptx_source(pptx_path: Path, repo_root: Path) -> SourceDocument:
    presentation = Presentation(str(pptx_path))
    slide_texts: list[str] = []
    for slide in presentation.slides:
        texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                texts.append(shape.text.strip())
        if texts:
            slide_texts.append("\n".join(texts))

    title = pptx_path.stem
    if slide_texts:
        first_line = next((line.strip() for line in slide_texts[0].splitlines() if line.strip()), "")
        if first_line:
            title = first_line

    return SourceDocument(
        relative_path=str(pptx_path.resolve().relative_to(repo_root)),
        file_name=pptx_path.name,
        file_extension=pptx_path.suffix.lower(),
        source_kind="pptx",
        quarter=_infer_quarter_from_path(pptx_path),
        title=title,
        extracted_text=_truncate_text("\n\n".join(slide_texts), 20000),
        file_size=pptx_path.stat().st_size,
        sha256=_sha256(pptx_path),
        metadata={
            "slide_count": len(presentation.slides),
            "extraction_status": "ok",
        },
    )


def _collect_ppt_source(ppt_path: Path, repo_root: Path) -> SourceDocument:
    return SourceDocument(
        relative_path=str(ppt_path.resolve().relative_to(repo_root)),
        file_name=ppt_path.name,
        file_extension=ppt_path.suffix.lower(),
        source_kind="ppt",
        quarter=_infer_quarter_from_path(ppt_path),
        title=ppt_path.stem,
        extracted_text="",
        file_size=ppt_path.stat().st_size,
        sha256=_sha256(ppt_path),
        metadata={
            "extraction_status": "unsupported_ppt",
        },
    )


def _collect_pdf_source(pdf_path: Path, repo_root: Path) -> SourceDocument:
    text_parts: list[str] = []
    page_count = 0
    with fitz.open(pdf_path) as document:
        page_count = len(document)
        for page in document:
            text_parts.append(page.get_text("text"))

    return SourceDocument(
        relative_path=str(pdf_path.resolve().relative_to(repo_root)),
        file_name=pdf_path.name,
        file_extension=pdf_path.suffix.lower(),
        source_kind="pdf",
        quarter=_infer_quarter_from_path(pdf_path),
        title=pdf_path.stem,
        extracted_text=_truncate_text("\n".join(text_parts), 20000),
        file_size=pdf_path.stat().st_size,
        sha256=_sha256(pdf_path),
        metadata={
            "page_count": page_count,
            "extraction_status": "ok",
        },
    )


def _build_error_source_document(path: Path, repo_root: Path, exc: Exception) -> SourceDocument:
    return SourceDocument(
        relative_path=str(path.resolve().relative_to(repo_root)),
        file_name=path.name,
        file_extension=path.suffix.lower(),
        source_kind=path.suffix.lower().lstrip("."),
        quarter=_infer_quarter_from_path(path),
        title=path.stem,
        extracted_text="",
        file_size=path.stat().st_size,
        sha256=_sha256(path),
        metadata={
            "extraction_status": "error",
            "error": str(exc),
        },
    )


def _collect_description_parts(row_values: dict[str, str]) -> list[tuple[str, str]]:
    parts: list[tuple[str, str]] = []
    for label, key in [("說明", "description"), ("RD Comment", "rd_comment"), ("Feature ID Detail", "feature_id_detail")]:
        value = row_values.get(key, "").strip()
        if value:
            parts.append((label, value))
    return parts


def _build_bug_record_from_mapping(
    *,
    row_values: dict[str, str],
    source_path: str,
    source_kind: str,
    source_sheet: str,
    source_row_number: int,
    quarter: str,
    source_priority: int,
    extra_description_parts: list[tuple[str, str]],
    extra_raw_fields: dict[str, Any],
) -> BugRecord:
    bug_id = _pick_first_value(row_values, "bug_id", "feature_id")
    project = _pick_first_value(row_values, "project")
    component = _pick_first_value(row_values, "component")
    status = _pick_first_value(row_values, "status", "closed_time")
    subject = _pick_first_value(row_values, "subject", "description")

    description_parts = extra_description_parts[:]
    if not description_parts:
        for label in ["description", "rd_comment", "feature_id_detail"]:
            value = row_values.get(label, "")
            if value:
                description_parts.append((label, value))
    description = _join_labeled_text(description_parts)
    if not description:
        description = _pick_first_value(row_values, "description", "rd_comment", "feature_id_detail")

    dedupe_key = _build_dedupe_key(bug_id=bug_id, project=project, subject=subject, mapped_values=row_values, source_path=source_path, source_sheet=source_sheet, source_row_number=source_row_number)
    period_rank = _quarter_rank(quarter or source_sheet or source_path)
    record_score = _score_record(
        bug_id=bug_id,
        project=project,
        component=component,
        status=status,
        subject=subject,
        description=description,
        period_rank=period_rank,
        source_priority=source_priority,
    )

    raw_json = {
        "source_path": source_path,
        "source_kind": source_kind,
        "source_sheet": source_sheet,
        "source_row_number": source_row_number,
        "quarter": quarter,
        "fields": row_values,
        "extra": extra_raw_fields,
    }

    return BugRecord(
        bug_id=bug_id,
        dedupe_key=dedupe_key,
        project=project,
        component=component,
        status=status,
        subject=subject,
        description=description,
        quarter=quarter,
        source_sheet=source_sheet,
        source_path=source_path,
        source_kind=source_kind,
        source_row_number=source_row_number,
        period_rank=period_rank,
        source_priority=source_priority,
        record_score=record_score,
        raw_json=raw_json,
    )


def _row_to_canonical_values(row: dict[int, Any], canonical_headers: dict[int, str]) -> dict[str, str]:
    values: dict[str, str] = {}
    for index, header_name in canonical_headers.items():
        if not header_name:
            continue
        value = row.get(index, "")
        if value in (None, ""):
            continue
        canonical_key = _canonical_key_for_header(header_name)
        if canonical_key:
            values[canonical_key] = str(value).strip()
    return values


def _row_to_mapping(row: tuple[Any, ...], header_map: dict[int, str]) -> dict[int, Any]:
    return {
        index: value
        for index, value in enumerate(row)
        if index in header_map and value not in (None, "")
    }


def _find_issue_sheet_header(sheet: Any) -> tuple[int | None, dict[int, str], dict[int, str]]:
    for row_number, row in enumerate(sheet.iter_rows(max_row=20, values_only=True), start=1):
        header_map = {
            index: _normalize_header(value)
            for index, value in enumerate(row)
            if _normalize_header(value)
        }
        if not header_map:
            continue
        matched_groups = {
            group_name
            for group_name, aliases in HEADER_GROUPS.items()
            if any(_header_matches(header_name, aliases) for header_name in header_map.values())
        }
        if (
            "project" in matched_groups
            and "subject" in matched_groups
            and ("bug_id" in matched_groups or "feature_id" in matched_groups)
            and any(group in matched_groups for group in {"component", "author", "assignee", "status"})
        ):
            return row_number, header_map, header_map
    return None, {}, {}


def _header_matches(header_name: str, aliases: list[str]) -> bool:
    normalized_aliases = [_normalize_header(alias) for alias in aliases]
    return any(alias == header_name or alias in header_name for alias in normalized_aliases)


def _canonical_key_for_header(header_name: str) -> str:
    normalized = _normalize_header(header_name)
    for canonical_key, aliases in [
        ("bug_id", ["bug id", "#", "issue id", "id"]),
        ("project", ["project", "案件"]),
        ("component", ["component"]),
        ("status", ["status"]),
        ("subject", ["subject", "title"]),
        ("description", ["說明", "description", "detail"]),
        ("rd_comment", ["rd comment", "rd command", "r&d solution for fix", "solution(root cause)", "solution (root cause)"]),
        ("author", ["author"]),
        ("assignee", ["assignee"]),
        ("feature_id_detail", ["feature id detail"]),
        ("feature_id", ["feature id"]),
        ("created_time", ["created time", "created"]),
        ("closed_time", ["closed"]),
        ("tracker", ["tracker"]),
        ("severity", ["severity"]),
        ("error_type", ["error type"]),
        ("issue_finder", ["issue finder"]),
        ("hw_version", ["hw version"]),
        ("fw_version", ["fw version"]),
        ("counts", ["counts"]),
        ("updated", ["updated"]),
        ("percent_done", ["% done", "percent done"]),
    ]:
        normalized_aliases = [_normalize_header(alias) for alias in aliases]
        if any(alias == normalized or alias in normalized for alias in normalized_aliases):
            return canonical_key
    return normalized.replace(" ", "_")


def _pick_first_value(values: dict[str, str], *keys: str) -> str:
    for key in keys:
        value = values.get(key, "")
        if value:
            return value
    return ""


def _build_dedupe_key(
    *,
    bug_id: str,
    project: str,
    subject: str,
    mapped_values: dict[str, str],
    source_path: str,
    source_sheet: str,
    source_row_number: int,
) -> str:
    if bug_id:
        return f"bug:{bug_id}"
    feature_id = mapped_values.get("feature_id", "")
    if feature_id:
        return f"feature:{project}:{feature_id}"
    if subject:
        return f"subject:{project}:{subject}"
    return f"row:{source_path}:{source_sheet}:{source_row_number}"


def _is_header_or_summary_row(record: BugRecord) -> bool:
    header_tokens = {
        "#",
        "bug id",
        "issue id",
        "id",
        "project",
        "component",
        "status",
        "subject",
        "feature id",
        "feature id detail",
    }
    core_values = [
        record.bug_id.strip().lower(),
        record.project.strip().lower(),
        record.component.strip().lower(),
        record.status.strip().lower(),
        record.subject.strip().lower(),
    ]
    header_hits = sum(value in header_tokens for value in core_values)
    if header_hits >= 3:
        return True
    if record.bug_id.strip() in {"#", "bug id", "issue id"}:
        return True
    if record.project.strip().lower() == "project" and record.subject.strip().lower() == "subject":
        return True
    return False


def _join_labeled_text(parts: list[tuple[str, str]]) -> str:
    lines: list[str] = []
    for label, text in parts:
        cleaned = str(text or "").strip()
        if not cleaned:
            continue
        lines.append(f"{label}: {cleaned}")
    return "\n".join(lines)


def _score_record(
    *,
    bug_id: str,
    project: str,
    component: str,
    status: str,
    subject: str,
    description: str,
    period_rank: int,
    source_priority: int,
) -> int:
    completeness = sum(1 for value in [bug_id, project, component, status, subject, description] if value)
    text_length = len(subject) + len(description)
    return period_rank * 100000 + source_priority * 1000 + completeness * 100 + min(text_length, 999)


def _infer_quarter_from_path(path: Path) -> str:
    parts = [path.stem, *path.parts]
    return _infer_quarter_from_text(" ".join(parts))


def _infer_quarter_from_text(text: str) -> str:
    matches = re.findall(r"(20\d{2})\s*[_ ]?q\s*([1-4])", text, flags=re.IGNORECASE)
    if not matches:
        return ""
    year, quarter = max(((int(year), int(q)) for year, q in matches), key=lambda item: (item[0], item[1]))
    return f"{year}Q{quarter}"


def _quarter_rank(text: str) -> int:
    quarter = _infer_quarter_from_text(text)
    if not quarter:
        return 0
    match = re.fullmatch(r"(20\d{2})Q([1-4])", quarter)
    if not match:
        return 0
    year = int(match.group(1))
    quarter_number = int(match.group(2))
    return year * 10 + quarter_number


def _normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    if not text:
        return ""
    text = text.replace("\u3000", " ")
    text = text.replace("_", " ")
    text = text.replace("\n", " ")
    text = re.sub(r"[\t\r]+", " ", text)
    text = re.sub(r"[()，,:;/\\\-]+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _collect_csv_source(csv_path: Path, repo_root: Path) -> tuple[SourceDocument, list[BugRecord], dict[str, Any]]:
    rows: list[dict[str, str]] = []
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        rows = [row for row in reader]

    relative_path = str(csv_path.resolve().relative_to(repo_root))
    document = SourceDocument(
        relative_path=relative_path,
        file_name=csv_path.name,
        file_extension=csv_path.suffix.lower(),
        source_kind="csv",
        quarter=_infer_quarter_from_text(csv_path.stem),
        title="Merged Redmine bug summary",
        extracted_text=f"CSV source with {len(rows)} rows.",
        file_size=csv_path.stat().st_size,
        sha256=_sha256(csv_path),
        metadata={
            "row_count": len(rows),
            "columns": list(rows[0].keys()) if rows else [],
        },
    )

    bug_records = []
    for row_number, row in enumerate(rows, start=2):
        mapped_values = {
            _canonical_key_for_header(key): str(value or "").strip()
            for key, value in row.items()
            if key
        }
        quarter = mapped_values.get("quarter", "")
        bug_records.append(
            _build_bug_record_from_mapping(
                row_values=mapped_values,
                source_path=relative_path,
                source_kind="csv",
                source_sheet=mapped_values.get("source_sheet", ""),
                source_row_number=row_number,
                quarter=quarter,
                source_priority=SOURCE_KIND_PRIORITY["csv"],
                extra_description_parts=[],
                extra_raw_fields={
                    "source_path": relative_path,
                    "source_kind": "csv",
                    "source_row_number": row_number,
                },
            )
        )

    return document, bug_records, {"row_count": len(rows)}


def _summarize_xlsx_document(sheet_stats: list[dict[str, Any]]) -> str:
    if not sheet_stats:
        return "Workbook contained no issue-like sheets."
    parts = [f"{item['sheet_name']} ({item['row_count']} rows)" for item in sheet_stats[:8]]
    return f"Workbook with {len(sheet_stats)} issue-like sheets: " + "; ".join(parts)


def _truncate_text(text: str, max_chars: int) -> str:
    if len(text) <= max_chars:
        return text
    return text[: max_chars - 3] + "..."


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


if __name__ == "__main__":
    main()
